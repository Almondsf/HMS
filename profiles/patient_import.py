'''
Import patients from a spreadsheet (CSV or Excel), e.g. when a clinic moves
from paper or Excel. Every row is checked with the same rules as registering
a patient by hand. A preview (dry run) lists problems by row and column
before anything is saved.
'''

import csv
import io
import re
from datetime import date, datetime

from django.db import transaction

from audit.recorder import Category, record

from .models import PatientProfile
from .serializers import PatientSerializer

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024

# Column -> friendly name, in template order
COLUMNS = {
    'first_name': 'First name',
    'last_name': 'Last name',
    'gender': 'Gender',
    'date_of_birth': 'Date of birth',
    'phone_number': 'Phone',
    'email': 'Email',
    'address': 'Address',
    'emergency_contact': 'Emergency contact',
    'emergency_phone': 'Emergency phone',
    'blood_group': 'Blood group',
    'genotype': 'Genotype',
    'allergies': 'Allergies',
    'medical_history': 'Medical history',
    'mrn': 'Existing record number',
}
REQUIRED = ('first_name', 'last_name')

# Header spellings people use, compared without spaces, case or punctuation
ALIASES = {
    **{re.sub(r'[^a-z]', '', label.lower()): field for field, label in COLUMNS.items()},
    **{field.replace('_', ''): field for field in COLUMNS},
    'givenname': 'first_name', 'forename': 'first_name', 'othernames': 'first_name', 'firstnames': 'first_name',
    'surname': 'last_name', 'familyname': 'last_name',
    'sex': 'gender',
    'dob': 'date_of_birth', 'birthdate': 'date_of_birth', 'dateofbirthddmmyyyy': 'date_of_birth',
    'phonenumber': 'phone_number', 'mobile': 'phone_number', 'mobilenumber': 'phone_number',
    'telephone': 'phone_number',
    'gsm': 'phone_number', 'tel': 'phone_number',
    'emailaddress': 'email', 'homeaddress': 'address',
    'nextofkin': 'emergency_contact', 'nextofkinphone': 'emergency_phone', 'nokphone': 'emergency_phone',
    'bloodgroup': 'blood_group', 'bloodtype': 'blood_group',
    'medicalhistory': 'medical_history', 'history': 'medical_history',
    'mrn': 'mrn', 'hospitalnumber': 'mrn', 'cardnumber': 'mrn', 'cardno': 'mrn', 'recordnumber': 'mrn',
    'foldernumber': 'mrn', 'folderno': 'mrn', 'patientnumber': 'mrn', 'patientid': 'mrn',
}  # fmt: skip

DATE_FORMATS = ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y')
GENDERS = {'m': 'male', 'male': 'male', 'f': 'female', 'female': 'female', 'o': 'other', 'other': 'other'}


class ImportFileError(Exception):
    '''The file itself can't be used (wrong type, too big, no usable columns).'''


def template_csv():
    return '﻿' + ','.join(COLUMNS.values()) + '\r\n'


def _text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores numbers as floats: 8031234567.0
    return str(value).strip()


def _read_rows(upload):
    '''[(spreadsheet row number, [cells])], header first.'''
    if upload.size > MAX_BYTES:
        raise ImportFileError('The file is larger than 5 MB. Split it into smaller files.')
    name = upload.name.lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise ImportFileError("This Excel file can't be read. Save it again as .xlsx or .csv.") from exc
        rows = [[_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    elif name.endswith('.csv'):
        raw = upload.read()
        for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):  # latin-1 accepts any bytes
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]
    else:
        raise ImportFileError('Upload a .csv or .xlsx file.')
    return [(number, row) for number, row in enumerate(rows, start=1) if any(row)]


def _map_header(header):
    mapping, unknown = {}, []
    for index, title in enumerate(header):
        field = ALIASES.get(re.sub(r'[^a-z]', '', title.lower()))
        if field and field not in mapping.values():
            mapping[index] = field
        elif title:
            unknown.append(title)
    missing = [COLUMNS[field] for field in REQUIRED if field not in mapping.values()]
    if missing:
        raise ImportFileError(
            f'Missing column(s): {", ".join(missing)}. The first row must hold the column names; '
            'download the template to see them.'
        )
    return mapping, unknown


def _date(value):
    for pattern in DATE_FORMATS:
        try:
            return datetime.strptime(value, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError('Use a date like 31/12/1990.')


def _phone(value):
    digits = re.sub(r'[\s\-().]', '', value)
    if re.fullmatch(r'[789]\d{9}', digits):
        digits = '0' + digits  # Excel drops the leading 0 of numbers like 0803...
    return digits


def _row_data(values):
    '''Spreadsheet values -> the patient API's input, plus problems found while converting.'''
    problems = {}
    user = {'first_name': values.get('first_name', ''), 'last_name': values.get('last_name', '')}
    for field in ('address', 'emergency_contact', 'email'):
        user[field] = values.get(field) or None
    if values.get('gender'):
        user['gender'] = GENDERS.get(values['gender'].lower())
        if user['gender'] is None:
            problems['gender'] = 'Use Male, Female or Other.'
    if values.get('date_of_birth'):
        try:
            user['date_of_birth'] = _date(values['date_of_birth'])
        except ValueError as exc:
            problems['date_of_birth'] = str(exc)
    for field in ('phone_number', 'emergency_phone'):
        if values.get(field):
            user[field] = _phone(values[field])
    data = {
        'user': user,
        'blood_group': values.get('blood_group', '').upper().replace(' ', '') or None,
        'genotype': values.get('genotype', '').upper().replace(' ', '') or None,
        'allergies': values.get('allergies', ''),
        'medical_history': values.get('medical_history', ''),
    }
    if values.get('mrn'):
        data['mrn'] = values['mrn']
    return data, problems


def _flatten(errors):
    '''{"user": {"email": ["..."]}} -> [("email", "...")]'''
    flat = []
    for field, value in errors.items():
        if isinstance(value, dict):
            flat += _flatten(value)
        else:
            messages = value if isinstance(value, list) else [value]
            flat += [(field, str(message)) for message in messages]
    return flat


def _preview(number, serializer):
    user = serializer.validated_data['user']
    return {
        'row': number,
        'name': f"{user['first_name']} {user['last_name']}",
        'gender': user.get('gender') or '',
        'date_of_birth': user.get('date_of_birth'),
        'phone_number': user.get('phone_number') or '',
        'mrn': serializer.validated_data.get('mrn') or '(new number)',
    }


def run_import(upload, *, commit=False, skip_invalid=False):
    '''
    Check every row; with commit, save the patients (all of them, or only
    the valid ones with skip_invalid). Nothing is saved if any row has a
    problem and skip_invalid is off.
    '''
    rows = _read_rows(upload)
    if not rows:
        raise ImportFileError('The file is empty.')
    (_, header), body = rows[0], rows[1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f'The file has {len(body)} patients; import at most {MAX_ROWS} at a time.')
    mapping, unknown = _map_header(header)

    existing = {
        (first.lower(), last.lower(), str(born)): mrn
        for first, last, born, mrn in PatientProfile.objects.filter(user__date_of_birth__isnull=False).values_list(
            'user__first_name', 'user__last_name', 'user__date_of_birth', 'mrn'
        )
    }
    seen_mrns, seen_emails = {}, {}
    errors, valid = [], []
    for number, cells in body:
        values = {field: (cells[index] if index < len(cells) else '') for index, field in mapping.items()}
        data, problems = _row_data(values)
        serializer = PatientSerializer(data=data)
        if not serializer.is_valid():
            for field, message in _flatten(serializer.errors):
                problems.setdefault(field, message)

        mrn, email, user = (data.get('mrn') or '').upper(), (data['user'].get('email') or '').lower(), data['user']
        if mrn and mrn in seen_mrns:
            problems.setdefault('mrn', f'Same record number as row {seen_mrns[mrn]}.')
        if email and email in seen_emails:
            problems.setdefault('email', f'Same email as row {seen_emails[email]}.')
        twin = existing.get((user['first_name'].lower(), user['last_name'].lower(), str(user.get('date_of_birth'))))
        if twin:
            problems.setdefault('first_name', f'Looks like {twin}, already registered (same name and date of birth).')
        if mrn:
            seen_mrns.setdefault(mrn, number)
        if email:
            seen_emails.setdefault(email, number)

        if problems:
            errors += [
                {'row': number, 'column': COLUMNS.get(field, field.replace('_', ' ').capitalize()), 'message': message}
                for field, message in problems.items()
            ]
        else:
            valid.append((number, serializer))

    imported = 0
    if commit and valid and (skip_invalid or not errors):
        with transaction.atomic():
            for _, serializer in valid:
                serializer.save()
            imported = len(valid)
            record(
                'patients.imported',
                Category.DOMAIN,
                metadata={'imported': imported, 'skipped_rows': len(body) - imported},
            )

    return {
        'total_rows': len(body),
        'valid_rows': len(valid),
        'rows_with_errors': len({error['row'] for error in errors}),
        'imported': imported,
        'errors': errors[:500],
        'unknown_columns': unknown,
        'preview': [_preview(number, serializer) for number, serializer in valid[:10]],
    }
